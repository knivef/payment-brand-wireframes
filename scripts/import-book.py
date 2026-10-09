#!/usr/bin/env python3
"""
Re-extract data/brand-book.json from the Brand Book workbook.

  python3 scripts/import-book.py ~/Downloads/Brand_Book.xlsx

The workbook stays the thing you edit. This script keeps the repo in step with it.
Run scripts/build.py afterwards so the pages pick up any rate changes.

Needs openpyxl:  pip3 install openpyxl
"""

import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

BRAND_KEYS = ["category", "brand", "domain", "priority", "status", "type", "licence",
              "sells", "industries", "price_lists", "pitches", "channels", "website",
              "deck", "pricing", "notes"]

PRICE_KEYS = ["name", "setup", "monthly", "dep_c2b", "dep_b2b", "wd_b2c", "wd_b2b",
              "settle_swift", "settle_sepa", "crypto_dep", "crypto_wd", "exchange",
              "reserve", "min_ticket", "max_ticket", "limit_rule", "next_step", "kyc"]


def clean(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v).strip())
    return s or None


def main():
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is not installed. Run: pip3 install openpyxl")

    if len(sys.argv) < 2:
        sys.exit("Usage: python3 scripts/import-book.py <path to Brand_Book.xlsx>")
    path = sys.argv[1]
    if not os.path.isfile(path):
        sys.exit("File not found: " + path)

    wb = openpyxl.load_workbook(path, data_only=True)

    brands = []
    for r in wb["Brands"].iter_rows(min_row=2, values_only=True):
        row = [clean(c) for c in r]
        if not row or not row[1]:
            continue
        d = {k: (row[i] if i < len(row) else None) for i, k in enumerate(BRAND_KEYS)}
        for f in ("price_lists", "pitches", "industries", "channels"):
            d[f] = [x.strip() for x in d[f].split(",")] if d.get(f) else []
        brands.append(d)

    prices = []
    for r in wb["Price lists"].iter_rows(min_row=6, values_only=True):
        row = [clean(c) for c in r]
        if not row or not row[0]:
            continue
        prices.append({k: (row[i] if i < len(row) else None) for i, k in enumerate(PRICE_KEYS)})

    pitches = []
    for r in wb["Pitches"].iter_rows(min_row=3, values_only=True):
        row = [clean(c) for c in r]
        if not row or not row[0]:
            continue
        pitches.append({"name": row[0],
                        "sells": row[1] if len(row) > 1 else None,
                        "industry": row[2] if len(row) > 2 else None,
                        "line": row[3] if len(row) > 3 else None})

    out = os.path.join(ROOT, "data", "brand-book.json")
    old = {}
    if os.path.isfile(out):
        with open(out, encoding="utf-8") as f:
            old = json.load(f)

    with open(out, "w", encoding="utf-8") as f:
        json.dump({"brands": brands, "price_lists": prices, "pitches": pitches}, f, indent=2)

    print("brands {0} (was {1}) | price lists {2} | pitches {3}".format(
        len(brands), len(old.get("brands", [])), len(prices), len(pitches)))

    old_rates = {p["name"]: p for p in old.get("price_lists", [])}
    for p in prices:
        o = old_rates.get(p["name"])
        if o and o != p:
            changed = [k for k in p if p.get(k) != o.get(k)]
            print("  rate changed on '{0}': {1}".format(p["name"], ", ".join(changed)))

    print("\nNow run: python3 scripts/build.py")


if __name__ == "__main__":
    main()
